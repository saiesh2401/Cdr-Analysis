from docx import Document

templates = ["JIO Template.docx", "Airtel Template.docx", "VI Template.docx"]

# The target format with placeholders
new_subject = "Subject:- Reg provide information in case the case FIR No. {FIR_NO}, PS Special Cell ({EMAIL})-{ISP_NAME}"

print("--- Updating Subject Lines ---")
for t in templates:
    try:
        doc = Document(t)
        found = False
        for p in doc.paragraphs:
            if p.text.strip().startswith("Subject"):
                p.text = new_subject
                found = True
                print(f"Updated {t}")
                break
        
        if not found:
            print(f"Subject line not found in {t}")
        else:
            doc.save(t)
            
    except Exception as e:
        print(f"Error processing {t}: {e}")
